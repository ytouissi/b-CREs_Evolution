#!/usr/bin/env python3
import os
import shutil
from pathlib import Path
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill

def count_sequences(file_path):
    with open(file_path, 'r') as f:
        count = sum(1 for line in f if line.startswith('>'))
    return count

def main():
    tracks = {
        'cactus447way': os.path.join(os.getcwd(), 'Intermediate_Files', 'Cactus_447_MSAs_Ready'),
        'multiz470way': os.path.join(os.getcwd(), 'Intermediate_Files', 'Multiz470_MSAs_Ready'),
        'ucsc30way': os.path.join(os.getcwd(), 'Intermediate_Files', 'UCSC30_MSAs_Ready')
    }
    
    output_folder = 'Adaptify_MSAs'
    os.makedirs(output_folder, exist_ok=True)
    
    # Collect all files from all tracks
    all_files = defaultdict(dict)
    
    for track_name, folder_path in tracks.items():
        if not os.path.exists(folder_path):
            print(f"Warning: {folder_path} not found")
            continue
        
        for file in os.listdir(folder_path):
            if file.endswith('.fa'):
                file_path = os.path.join(folder_path, file)
                filename = Path(file).stem
                parts = filename.split('_')
                
                if len(parts) >= 3:
                    region_key = f"{parts[0]}_{parts[1]}_{parts[2]}"
                    seq_count = count_sequences(file_path)
                    
                    all_files[region_key][track_name] = {
                        'file': file_path,
                        'count': seq_count
                    }
    
    # Process each region
    excel_data = []
    
    for region_key in sorted(all_files.keys()):
        track_data = all_files[region_key]
        
        # If region exists in multiple tracks, select best
        if len(track_data) > 1:
            best_track = max(track_data.keys(), key=lambda t: track_data[t]['count'])
        else:
            best_track = list(track_data.keys())[0]
        
        best_data = track_data[best_track]
        src_file = best_data['file']
        dst_file = os.path.join(output_folder, os.path.basename(src_file))
        shutil.copy2(src_file, dst_file)
        
        # Prepare Excel entry
        parts = region_key.split('_')
        coordinates = f"{parts[0]}:{parts[1]}-{parts[2]}"
        
        excel_data.append({
            'Coordinates': coordinates,
            'SeqCount': best_data['count'],
            'Track': best_track
        })
    
    # Create Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'MSA Selection'
    
    headers = ['Coordinates', 'SeqCount', 'Track']
    ws.append(headers)
    
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    
    for row in excel_data:
        ws.append([row['Coordinates'], row['SeqCount'], row['Track']])
    
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 15
    
    wb.save('MSA_Selection_Report.xlsx')
    
    print(f"Total regions: {len(excel_data)}")
    print(f"Files copied to: {output_folder}")
    print(f"Report saved: MSA_Selection_Report.xlsx")

if __name__ == "__main__":
    main()